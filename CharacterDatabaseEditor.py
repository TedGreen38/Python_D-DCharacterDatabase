"""Display and edit ``CharacterDatabase.xlsx`` through a small Tkinter GUI.

The workbook is loaded into an in-memory SQLite table so rows can be listed,
added, and updated with SQL. Changes are written back to the workbook after
each add or update operation.
"""
# imports sqlite3 for sqlite operations and functionality
import sqlite3
from pathlib import Path

# imports tkinter for GUI functionality.
import tkinter as tk
from tkinter import ttk

try:
	# Keep the application usable as a clear command-line setup error if the
	# workbook-reading dependency has not been installed.
	from openpyxl import load_workbook
except ImportError as exc:
	raise SystemExit("Install openpyxl first: pip install openpyxl") from exc


class CharacterDatabaseModel:
	"""Load workbook data, expose SQL-backed operations, and persist changes."""

	def __init__(self, path):
		"""Open the workbook and copy its non-empty data rows into SQLite."""
		self.path = path
		self.workbook = load_workbook(path)
		self.sheet = self.workbook.active
		
		# Treat the first worksheet row as column labels; ignore trailing blank cells.
		self.headers = [cell.value for cell in self.sheet[1]]
		while self.headers and self.headers[-1] is None:
			self.headers.pop()
		if not self.headers or any(header is None for header in self.headers):
			raise ValueError("The first row must contain a column name in every used column.")

		# Internal SQL-safe column names let workbook headers contain arbitrary text.
		self.columns = [f"column_{index}" for index in range(len(self.headers))]
		self.sql_columns = ", ".join(f'"{column}"' for column in self.columns)
		self.placeholders = ", ".join("?" for _ in self.columns)
		self.database = sqlite3.connect(":memory:")
		self.database.execute(f"CREATE TABLE characters ({self.sql_columns})")
		
		# Skip fully blank worksheet rows and convert date-like values to text for SQL.
		for values in self.sheet.iter_rows(
			min_row=2, max_col=len(self.headers), values_only=True
		):
			if any(value is not None for value in values):
				self.database.execute(
					f"INSERT INTO characters ({self.sql_columns}) VALUES ({self.placeholders})",
					tuple(value.isoformat() if hasattr(value, "isoformat") else value for value in values),
				)
		self.database.commit()

	def get_rows(self):
		"""Return every row, including its SQLite row ID, in display order."""
		return self.database.execute(
			f"SELECT rowid, {self.sql_columns} FROM characters ORDER BY rowid"
		).fetchall()

	def get_character(self, row_id):
		"""Return the values for one character, or ``None`` if it is missing."""
		return self.database.execute(
			f"SELECT {self.sql_columns} FROM characters WHERE rowid = ?", (row_id,)
		).fetchone()

	def add_character(self, values):
		"""Insert a character and return its newly assigned row ID."""
		cursor = self.database.execute(
			f"INSERT INTO characters ({self.sql_columns}) VALUES ({self.placeholders})",
			values,
		)
		return cursor.lastrowid

	def update_character(self, row_id, values):
		"""Replace all column values for the character with the given row ID."""
		assignments = ", ".join(f'"{column}" = ?' for column in self.columns)
		self.database.execute(
			f"UPDATE characters SET {assignments} WHERE rowid = ?", (*values, row_id)
		)

	def save(self):
		"""Write the current SQL rows to the worksheet and save the workbook."""
		rows = self.database.execute(
			f"SELECT {self.sql_columns} FROM characters ORDER BY rowid"
		).fetchall()
		# Rewrite data rows in order, then clear any leftover cells from removed rows.
		for row_number, values in enumerate(rows, start=2):
			for column_number, value in enumerate(values, start=1):
				self.sheet.cell(row=row_number, column=column_number, value=value)
		for row_number in range(len(rows) + 2, self.sheet.max_row + 1):
			for column_number in range(1, len(self.headers) + 1):
				self.sheet.cell(row=row_number, column=column_number).value = None
		self.workbook.save(self.path)

	def close(self):
		"""Release the in-memory SQLite connection."""
		self.database.close()


class CharacterDatabaseViewModel:
	"""Track the selected character and coordinate GUI actions with the model."""

	def __init__(self, model):
		self.model = model
		self.selected_id = None

    # Forward the model's current rows for display in the table.
	def get_rows(self):
		return self.model.get_rows()

    # Remember a row selection and fetch that character's values.
	def select_character(self, row_id):
		self.selected_id = row_id
		return self.model.get_character(row_id)

    # Clear the current row selection.
	def clear_selection(self):
		self.selected_id = None

    # Add a character, select it, and immediately persist the workbook.
	def add_character(self, values):
		self.selected_id = self.model.add_character(values)
		self.model.save()
		return self.selected_id

    # Update and save the selected character; return false if none is selected.
	def update_character(self, values):
		if self.selected_id is None:
			return False
		self.model.update_character(self.selected_id, values)
		self.model.save()
		return True

# Build the editor window, wire up its controls, and run the event loop.
def main():
	# Resolve the workbook beside this script so the program works from any cwd.
	path = Path(__file__).with_name("CharacterDatabase.xlsx")
	if not path.is_file():
		print(f"Workbook not found: {path}")
		return

	try:
		model = CharacterDatabaseModel(path)
	except ValueError as exc:
		print(exc)
		return
	view_model = CharacterDatabaseViewModel(model)
	headers = model.headers
	columns = model.columns
	
	# Configure the main window and a two-column layout for the table and form.
	root = tk.Tk()
	root.title("Character Database Editor")
	root.geometry("1100x650")
	
	main_frame = ttk.Frame(root, padding=10)
	main_frame.pack(fill="both", expand=True)
	main_frame.columnconfigure(0, weight=3)
	main_frame.columnconfigure(1, weight=2)
	main_frame.rowconfigure(0, weight=1)

	# The left pane shows workbook rows; selecting a row loads it into the form.
	table_frame = ttk.Frame(main_frame)
	table_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
	table_frame.rowconfigure(0, weight=1)
	table_frame.columnconfigure(0, weight=1)
	tree = ttk.Treeview(table_frame, columns=columns, show="headings")
	tree.grid(row=0, column=0, sticky="nsew")
	vertical_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
	vertical_scroll.grid(row=0, column=1, sticky="ns")
	horizontal_scroll = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
	horizontal_scroll.grid(row=1, column=0, sticky="ew")
	tree.configure(yscrollcommand=vertical_scroll.set, xscrollcommand=horizontal_scroll.set)
	tree.heading("#0", text="ID")
	for column, header in zip(columns, headers):
		tree.heading(column, text=str(header))
		tree.column(column, width=120, minwidth=70, stretch=True)

	# The right pane provides one editable text field for each workbook column.
	form = ttk.LabelFrame(main_frame, text="Selected character", padding=10)
	form.grid(row=0, column=1, sticky="nsew")
	form.rowconfigure(0, weight=1)
	form.columnconfigure(0, weight=1)
	canvas = tk.Canvas(form, highlightthickness=0)
	canvas.grid(row=0, column=0, sticky="nsew")
	form_scroll = ttk.Scrollbar(form, orient="vertical", command=canvas.yview)
	form_scroll.grid(row=0, column=1, sticky="ns")
	canvas.configure(yscrollcommand=form_scroll.set)
	fields_frame = ttk.Frame(canvas)
	canvas_window = canvas.create_window((0, 0), window=fields_frame, anchor="nw")
	fields_frame.bind("<Configure>", lambda _event: canvas.configure(scrollregion=canvas.bbox("all")))
	canvas.bind("<Configure>", lambda event: canvas.itemconfigure(canvas_window, width=event.width))

	# Create form labels and entry fields from the workbook's header row.
	field_values = []
	for row_number, header in enumerate(headers):
		ttk.Label(fields_frame, text=str(header)).grid(row=row_number, column=0, sticky="w", pady=3, padx=(0, 8))
		value = tk.StringVar()
		ttk.Entry(fields_frame, textvariable=value).grid(row=row_number, column=1, sticky="ew", pady=3)
		field_values.append(value)
	fields_frame.columnconfigure(1, weight=1)

	status = tk.StringVar(value="Select a row or enter values to add a character.")
	ttk.Label(form, textvariable=status, wraplength=300).grid(
		row=1, column=0, columnspan=2, sticky="w", pady=(10, 0)
	)
	# Refresh the table while optionally restoring the selection after an edit.
	def refresh_rows(select_id=None):
		tree.delete(*tree.get_children())
		for row in view_model.get_rows():
			row_id, *values = row
			iid = str(row_id)
			tree.insert("", "end", iid=iid, text=iid, values=values)
		if select_id is not None and tree.exists(str(select_id)):
			tree.selection_set(str(select_id))
			tree.focus(str(select_id))

	# Clear form contents and selection to prepare for adding a new character.
	def clear_fields():
		for value in field_values:
			value.set("")
		tree.selection_remove(tree.selection())
		view_model.clear_selection()
		status.set("Enter values to add a character, or select a row to edit it.")

	# Populate the form whenever the user selects a character in the table.
	def select_row(_event=None):
		selection = tree.selection()
		if not selection:
			return
		
		row_id = int(selection[0])
		row = view_model.select_character(row_id)
		if row is not None:
			for field, value in zip(field_values, row):
				field.set("" if value is None else str(value))
			status.set(f"Editing row {row_id}.")

	# Add the form contents as a new row, save, and select the inserted row.
	def add_row():
		values = [field.get() for field in field_values]
		row_id = view_model.add_character(values)
		refresh_rows(row_id)
		status.set(f"Saved changes to {path.name}.")

	# Save updated form contents to the currently selected row.
	def update_row():
		if not view_model.update_character([field.get() for field in field_values]):
			status.set("Select a row before updating it.")
			return
		refresh_rows(view_model.selected_id)
		status.set(f"Saved changes to {path.name}.")

	# Place add, update, and clear actions beneath the editable form.
	button_frame = ttk.Frame(form)
	button_frame.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(10, 0))
	ttk.Button(button_frame, text="Add", command=add_row).pack(side="left", padx=(0, 5))
	ttk.Button(button_frame, text="Update selected", command=update_row).pack(side="left", padx=5)
	ttk.Button(button_frame, text="Clear", command=clear_fields).pack(side="left", padx=5)
	tree.bind("<<TreeviewSelect>>", select_row)
	refresh_rows()

	# Close the SQL connection before destroying the application window.
	def close():
		model.close()
		root.destroy()

	root.protocol("WM_DELETE_WINDOW", close)
	root.mainloop()


if __name__ == "__main__":
	main()
